#!/usr/bin/env python3
"""Turn a .zip of images into a .pptx with one image per slide.

Usage: python3 zip2pptx.py [--progress] images.zip [output.pptx]
  --progress  print "PROGRESS <done> <total>" after each slide (used by zip2pptx-progress.js)
Requires: pip install python-pptx
"""
import io
import re
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)  # 16:9 (13.333in x 7.5in)


def natural_key(name):
    # Sort "img2" before "img10".
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", name)]


def main():
    args = sys.argv[1:]
    progress = "--progress" in args
    args = [a for a in args if a != "--progress"]
    if not args:
        sys.exit(__doc__)
    zip_path = Path(args[0])
    out_path = Path(args[1]) if len(args) > 1 else zip_path.with_suffix(".pptx")

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    with zipfile.ZipFile(zip_path) as zf:
        names = sorted(
            (n for n in zf.namelist()
             if Path(n).suffix.lower() in IMAGE_EXTS
             and not n.startswith("__MACOSX/")
             and not Path(n).name.startswith(".")),
            key=natural_key,
        )
        if not names:
            sys.exit(f"No images found in {zip_path}")

        for i, name in enumerate(names, 1):
            data = zf.read(name)
            with Image.open(io.BytesIO(data)) as img:
                w, h = img.size
                # PowerPoint can't embed TIFF/BMP reliably; convert those to PNG.
                if img.format not in ("JPEG", "PNG", "GIF"):
                    buf = io.BytesIO()
                    img.save(buf, "PNG")
                    data = buf.getvalue()

            # Fit inside the slide, preserving aspect ratio, centered.
            scale = min(SLIDE_W / w, SLIDE_H / h)
            pw, ph = int(w * scale), int(h * scale)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                io.BytesIO(data), (SLIDE_W - pw) // 2, (SLIDE_H - ph) // 2, pw, ph
            )
            if progress:
                print(f"PROGRESS {i} {len(names)}", flush=True)

    if progress:
        print("SAVING", flush=True)

    prs.save(out_path)
    print(f"Wrote {len(names)} slides to {out_path}")


if __name__ == "__main__":
    main()
